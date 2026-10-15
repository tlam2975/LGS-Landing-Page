#!/usr/bin/env python3
"""Rebuild mb-promotion/data.json from the monthly MB promotion workbook.

Usage (from the repo root):
    python3 mb-promotion/build_data.py "1. MB Internal Promotion_Tháng 9.xlsx"
Requires: pip install openpyxl
"""
import json, re, sys, datetime
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parent.parent
src = Path(sys.argv[1]) if len(sys.argv) > 1 else next(root.glob("1. MB Internal Promotion*.xlsx"))
ws = openpyxl.load_workbook(src, data_only=True).active

# Product images: reuse the maps already maintained in assets/catalog.js
# (local files in assets/images win over LG-hosted URLs).
js = (root / "assets/catalog.js").read_text(encoding="utf-8")
images = {}
for m in re.finditer(r'^\s*"?([A-Za-z0-9\-]+)"?:\s*"(https?://[^"]+)"', js, re.M):
    images[m.group(1)] = m.group(2)
for m in re.finditer(r'^\s*"?([A-Za-z0-9\-]+)"?:\s*"(\.\./assets/images/[^"]+)"', js, re.M):
    images[m.group(1)] = m.group(2)

def num(v):
    return None if v in (None, "") else int(round(float(v)))

def clean(v):
    return " ".join(str(v).split()) if v else ""

products = []
for r in ws.iter_rows(min_row=9, max_col=17, values_only=True):
    if not r[0] or not r[1]:
        continue
    model = (r[2] or "").strip()
    base = model.split(".")[0]
    products.append({
        "cat": r[0].strip(),
        "name": clean(r[1]),
        "model": model,
        "staffVoucher": num(r[3]),
        "outright": num(r[5]),
        "monthly": {"12": num(r[6]), "24": num(r[7]), "36": num(r[8])},
        "yearly": {"1": num(r[11]), "2": num(r[12]), "3": num(r[13])},
        "image": images.get(base, ""),
        "warranty": clean(r[16]),  # col Q "Gia hạn bảo hành" (TV only)
    })

validity = clean(ws["P6"].value)
m = re.search(r"(\d{2})/(\d{2})/(\d{4})", validity)
valid_until = f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else None

customer = [{"minOrder": num(ws[f"T{r}"].value), "gift": clean(ws[f"U{r}"].value)} for r in (9, 10)]

kiosk = []
for row in (17, 18):
    gift = clean(ws[f"U{row}"].value)
    g = re.match(r"(.+?)\s*-\s*(\S+)\s*-\s*giá\s*(\S+)", gift)
    kiosk.append({
        "minOrder": num(ws[f"T{row}"].value),
        "gift": gift,
        "giftName": g.group(1) if g else gift,
        "giftModel": g.group(2) if g else "",
        "giftPrice": g.group(3) if g else "",
        "image": images.get(g.group(2).split(".")[0], "") if g else "",
    })

data = {
    "title": clean(ws["A1"].value),
    "source": src.name,
    "generated": datetime.date.today().isoformat(),
    "validUntil": valid_until,  # edit here if the offer is extended
    "validityText": validity,
    "staffIntro": clean(ws["A4"].value),
    "staffNote": clean(ws["A5"].value),
    "customerIntro": clean(ws["S6"].value),
    "customerTiers": customer,
    "kioskIntro": clean(ws["S13"].value),
    "kioskTiers": kiosk,
    "notes": [clean(ws[f"S{r}"].value) for r in (21, 22, 23)],
    "products": products,
}
out = Path(__file__).with_name("data.json")
out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
print(f"{len(products)} products -> {out}")
